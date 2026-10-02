import json
import logging
from pathlib import Path
from typing import List, Dict, Any, Union
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configuração de logging integrada
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")


class FinancialReporter:
    """
    Engine de geração de relatórios e exportação multi-formato.
    Converte dados auditados em entregáveis estruturados (JSON) e planilhas
    de relatórios executivos (Excel) estilizadas com formatação condicional.
    """

    def __init__(self, output_dir: Union[str, Path] = "data/output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_to_json(self, records: List[Dict[str, Any]], filename: str = "dados_estruturados.json") -> Path:
        """Serializa os registros auditados num arquivo JSON legível e higienizado."""
        target_path = self.output_dir / filename
        try:
            with open(target_path, "w", encoding="utf-8") as f:
                json.dump(records, f, ensure_ascii=False, indent=4)
            logging.info(f"Relatório JSON gerado com sucesso: {target_path}")
            return target_path
        except Exception as e:
            logging.error(f"Erro ao exportar JSON para {target_path}: {e}")
            raise

    def export_to_excel(self, records: List[Dict[str, Any]], filename: str = "relatorio_auditoria.xlsx") -> Path:
        """
        Gera uma planilha Excel estilizada contendo os dados detalhados da auditoria
        e uma aba executiva com métricas agregadas de risco.
        """
        target_path = self.output_dir / filename
        if not records:
            logging.warning("Nenhum registro fornecido para exportação em Excel.")
            return target_path

        df_details = pd.DataFrame(records)

        # Mapeamento e reordenação das colunas para visualização corporativa
        column_mapping = {
            "line_number": "Linha Original",
            "id": "ID Transação",
            "date": "Data",
            "description": "Descrição",
            "amount": "Valor (R$)",
            "is_flagged": "Sinalizado",
            "anomalies_count": "Qtd Anomalias",
            "anomalies": "Inconsistências Identificadas",
            "severity": "Nível de Severidade"
        }
        
        existing_cols = [col for col in column_mapping.keys() if col in df_details.columns]
        df_details = df_details[existing_cols].rename(columns=column_mapping)

        # Apuração de KPIs para o Resumo Executivo
        total_records = len(records)
        flagged_records = sum(1 for r in records if r.get("is_flagged", False))
        critical_records = sum(1 for r in records if r.get("severity") == "CRITICAL")
        total_amount = sum(r.get("amount", 0.0) for r in records)
        risk_amount = sum(abs(r.get("amount", 0.0)) for r in records if r.get("is_flagged", False))

        df_summary = pd.DataFrame([
            {"Métrica Operacional": "Total de Transações Processadas", "Valor": total_records},
            {"Métrica Operacional": "Transações com Anomalias (Flagged)", "Valor": flagged_records},
            {"Métrica Operacional": "Inconformidades Críticas", "Valor": critical_records},
            {"Métrica Operacional": "Volume Monetário Total (R$)", "Valor": f"{total_amount:,.2f}"},
            {"Métrica Operacional": "Volume Monetário em Risco (R$)", "Valor": f"{risk_amount:,.2f}"}
        ])

        try:
            with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
                df_summary.to_excel(writer, sheet_name="Resumo_Executivo", index=False)
                df_details.to_excel(writer, sheet_name="Detalhamento_Auditoria", index=False)

                self._format_summary_sheet(writer.sheets["Resumo_Executivo"])
                self._format_details_sheet(writer.sheets["Detalhamento_Auditoria"], df_details)

            logging.info(f"Relatório Excel gerado com sucesso: {target_path}")
            return target_path
        except Exception as e:
            logging.error(f"Erro ao exportar Excel para {target_path}: {e}")
            raise

    def _format_summary_sheet(self, ws) -> None:
        """Aplica estilo corporativo na aba de resumo executivo."""
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border = Border(bottom=Side(style="thin", color="D9D9D9"))

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=2, max_col=2):
            for cell in row:
                cell.border = border

        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 25

    def _format_details_sheet(self, ws, df: pd.DataFrame) -> None:
        """Aplica estilização condicional por nível de severidade e ajuste dinâmico de colunas."""
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        # Paleta de cores para destaque condicional de severidade
        critical_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Vermelho/Laranja Suave
        warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # Amarelo Suave

        critical_font = Font(color="C00000", bold=True)
        warning_font = Font(color="B25900", bold=True)

        # Formatação do Cabeçalho
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        severity_col_idx = None
        for idx, col in enumerate(df.columns, 1):
            if col == "Nível de Severidade":
                severity_col_idx = idx
                break

        # Formatação por linha conforme a severidade apontada
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=len(df) + 1), 2):
            severity_val = str(ws.cell(row=row_idx, column=severity_col_idx).value) if severity_col_idx else ""
            
            fill_to_apply = None
            font_to_apply = None
            
            if severity_val == "CRITICAL":
                fill_to_apply = critical_fill
                font_to_apply = critical_font
            elif severity_val == "WARNING":
                fill_to_apply = warning_fill
                font_to_apply = warning_font

            for cell in row:
                if fill_to_apply and font_to_apply:
                    cell.fill = fill_to_apply
                    cell.font = font_to_apply

        # Auto-fit de largura das colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)